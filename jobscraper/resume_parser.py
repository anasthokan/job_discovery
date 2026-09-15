"""Extract structured profile data from a resume / CV."""

from __future__ import annotations

import io
import re
from typing import Any

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
SUPPORTED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "text/plain",
    "text/markdown",
    "application/octet-stream",
}

EMAIL_RE = re.compile(r"\b[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}\b", re.I)
PHONE_RE = re.compile(
    r"(?:(?:\+|00)\d{1,3}[\s.\-]?)?(?:\(?\d{2,4}\)?[\s.\-]?)?\d{3,5}[\s.\-]?\d{4,6}"
)
URL_RE = re.compile(r"https?://[^\s)>\]]+|www\.[^\s)>\]]+", re.I)
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[A-Za-z0-9\-_%/]+", re.I)
GITHUB_RE = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[A-Za-z0-9\-_.]+/?", re.I)
CITY_STATE_RE = re.compile(
    r"\b([A-Z][A-Za-z .'-]+),\s*([A-Z]{2}|[A-Z][A-Za-z]+(?:\s[A-Z][A-Za-z]+)?)\b"
)
YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")
GPA_RE = re.compile(r"\bGPA[:\s]*([0-4](?:\.\d{1,2})?)(?:\s*/\s*4(?:\.0+)?)?\b", re.I)

MONTH = (
    r"Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|"
    r"Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?"
)
DATE_TOKEN = rf"(?:{MONTH})\.?\s+\d{{4}}|\d{{1,2}}[/-]\d{{4}}|\d{{4}}"
DATE_RANGE_RE = re.compile(
    rf"(?P<start>{DATE_TOKEN})\s*(?:[-–—]|to)\s+(?P<end>Present|Current|Now|{DATE_TOKEN})",
    re.I,
)
DATE_RANGE_LINE_RE = re.compile(
    rf"^(?P<start>{DATE_TOKEN})\s*(?:[-–—]|to)\s+(?P<end>Present|Current|Now|{DATE_TOKEN})\s*$",
    re.I,
)

SECTION_ALIASES = {
    "summary": ("summary", "professional summary", "profile", "about", "objective", "career objective"),
    "experience": (
        "experience",
        "work experience",
        "professional experience",
        "employment",
        "employment history",
        "work history",
        "career history",
    ),
    "education": ("education", "academic background", "academics", "qualifications"),
    "skills": (
        "skills",
        "technical skills",
        "core skills",
        "key skills",
        "competencies",
        "technologies",
        "tech stack",
        "expertise",
    ),
    "projects": ("projects", "personal projects", "selected projects", "key projects"),
    "certifications": ("certifications", "certificates", "licenses", "licenses & certifications"),
    "languages": ("languages", "language skills"),
}

SECTION_LOOKUP = {
    alias: key
    for key, aliases in SECTION_ALIASES.items()
    for alias in aliases
}

TITLE_HINTS = (
    "engineer",
    "developer",
    "programmer",
    "manager",
    "analyst",
    "intern",
    "consultant",
    "designer",
    "architect",
    "lead",
    "director",
    "specialist",
    "coordinator",
    "administrator",
    "scientist",
    "researcher",
    "founder",
    "officer",
    "associate",
    "assistant",
    "head",
    "principal",
    "staff",
    "sre",
    "devops",
    "qa",
    "tester",
    "product",
    "owner",
    "scrum",
    "data",
    "ml",
    "ai",
    "ceo",
    "cto",
    "cfo",
    "coo",
    "vp",
    "president",
)

DEGREE_RE = re.compile(
    r"\b("
    r"Ph\.?D\.?|Doctor of Philosophy|"
    r"M\.?B\.?A\.?|Master of Business Administration|"
    r"M\.?S\.?c?\.?|M\.?A\.?|M\.?Tech\.?|M\.?E\.?|Master(?:'s)?(?: of [A-Za-z &]+)?"
    r"|B\.?S\.?c?\.?|B\.?A\.?|B\.?Tech\.?|B\.?E\.?|Bachelor(?:'s)?(?: of [A-Za-z &]+)?"
    r"|Associate(?:'s)?(?: Degree)?"
    r"|Diploma|High School|HSC|SSC|Intermediate"
    r")(?=\s|,|$)",
    re.I,
)

SKILL_CANONICAL = {
    "python": "Python",
    "javascript": "JavaScript",
    "typescript": "TypeScript",
    "java": "Java",
    "c#": "C#",
    "c++": "C++",
    "c": "C",
    "go": "Go",
    "golang": "Go",
    "rust": "Rust",
    "php": "PHP",
    "ruby": "Ruby",
    "kotlin": "Kotlin",
    "swift": "Swift",
    "scala": "Scala",
    "r": "R",
    "sql": "SQL",
    "mysql": "MySQL",
    "postgresql": "PostgreSQL",
    "postgres": "PostgreSQL",
    "mongodb": "MongoDB",
    "redis": "Redis",
    "sqlite": "SQLite",
    "oracle": "Oracle",
    "mssql": "SQL Server",
    "sql server": "SQL Server",
    "html": "HTML",
    "css": "CSS",
    "sass": "Sass",
    "react": "React",
    "react.js": "React",
    "reactjs": "React",
    "next.js": "Next.js",
    "nextjs": "Next.js",
    "vue": "Vue.js",
    "vue.js": "Vue.js",
    "angular": "Angular",
    "node": "Node.js",
    "node.js": "Node.js",
    "nodejs": "Node.js",
    "express": "Express",
    "fastapi": "FastAPI",
    "django": "Django",
    "flask": "Flask",
    "spring": "Spring",
    "spring boot": "Spring Boot",
    ".net": ".NET",
    "dotnet": ".NET",
    "asp.net": "ASP.NET",
    "laravel": "Laravel",
    "rails": "Ruby on Rails",
    "graphql": "GraphQL",
    "rest": "REST",
    "rest api": "REST",
    "aws": "AWS",
    "azure": "Azure",
    "gcp": "GCP",
    "google cloud": "GCP",
    "docker": "Docker",
    "kubernetes": "Kubernetes",
    "k8s": "Kubernetes",
    "terraform": "Terraform",
    "ansible": "Ansible",
    "linux": "Linux",
    "git": "Git",
    "github": "GitHub",
    "gitlab": "GitLab",
    "ci/cd": "CI/CD",
    "jenkins": "Jenkins",
    "pandas": "Pandas",
    "numpy": "NumPy",
    "scikit-learn": "scikit-learn",
    "sklearn": "scikit-learn",
    "tensorflow": "TensorFlow",
    "pytorch": "PyTorch",
    "nlp": "NLP",
    "machine learning": "Machine Learning",
    "deep learning": "Deep Learning",
    "data analysis": "Data Analysis",
    "power bi": "Power BI",
    "tableau": "Tableau",
    "excel": "Excel",
    "spark": "Apache Spark",
    "hadoop": "Hadoop",
    "airflow": "Airflow",
    "kafka": "Kafka",
    "elasticsearch": "Elasticsearch",
    "scrapy": "Scrapy",
    "selenium": "Selenium",
    "playwright": "Playwright",
    "pytest": "pytest",
    "junit": "JUnit",
    "jira": "Jira",
    "agile": "Agile",
    "scrum": "Scrum",
    "figma": "Figma",
    "tailwind": "Tailwind CSS",
    "bootstrap": "Bootstrap",
    "redux": "Redux",
    "webpack": "Webpack",
    "bash": "Bash",
    "powershell": "PowerShell",
    "firebase": "Firebase",
    "supabase": "Supabase",
    "s3": "S3",
    "lambda": "AWS Lambda",
    "ec2": "EC2",
    "snowflake": "Snowflake",
    "dbt": "dbt",
    "looker": "Looker",
    "huggingface": "Hugging Face",
    "openai": "OpenAI",
    "langchain": "LangChain",
}

SKILL_PATTERN = re.compile(
    r"\b("
    + "|".join(
        re.escape(key)
        for key in sorted(SKILL_CANONICAL, key=len, reverse=True)
        if len(key) >= 3
    )
    + r")\b",
    re.I,
)


class ResumeParseError(ValueError):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


def parse_resume_bytes(data: bytes, filename: str = "", content_type: str = "") -> dict[str, Any]:
    if not data:
        raise ResumeParseError("Empty file.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ResumeParseError("File is larger than 8 MB.", 413)

    ext = _extension(filename)
    if ext and ext not in SUPPORTED_EXTENSIONS:
        raise ResumeParseError(f"Unsupported file type '{ext}'. Use PDF, DOCX, or TXT.", 415)

    ctype = (content_type or "").split(";")[0].strip().lower()
    if ctype and ctype not in SUPPORTED_CONTENT_TYPES and not ext:
        raise ResumeParseError(f"Unsupported content type '{ctype}'.", 415)

    text, pages = _extract_text(data, ext, ctype)
    cleaned = _normalize_text(text)
    if len(cleaned.strip()) < 40:
        raise ResumeParseError("Could not extract enough text from this file.", 422)

    return parse_resume_text(cleaned, filename=filename, content_type=ctype, pages=pages)


def parse_resume_text(
    text: str,
    filename: str = "",
    content_type: str = "text/plain",
    pages: int | None = None,
) -> dict[str, Any]:
    cleaned = _normalize_text(text)
    lines = [line.strip() for line in cleaned.splitlines() if line.strip()]
    sections = _split_sections(lines)
    email = _first_match(EMAIL_RE, cleaned)
    phones = _unique(_normalize_phone(match) for match in PHONE_RE.findall(cleaned))
    phones = [phone for phone in phones if _looks_like_phone(phone)]
    linkedin = _first_match(LINKEDIN_RE, cleaned)
    github = _first_match(GITHUB_RE, cleaned)
    website = _pick_website(cleaned, linkedin, github)
    location = _extract_location(lines, sections)
    skills = _extract_skills(sections.get("skills", []), cleaned)
    return {
        "ok": True,
        "source": {
            "filename": filename or None,
            "content_type": content_type or None,
            "pages": pages,
            "char_count": len(cleaned),
        },
        "profile": {
            "full_name": _extract_name(lines, email, phones),
            "email": email,
            "phone": phones[0] if phones else None,
            "phones": phones,
            "location": location,
            "linkedin": _normalize_url(linkedin),
            "github": _normalize_url(github),
            "website": _normalize_url(website),
            "summary": _extract_summary(sections, lines),
        },
        "skills": skills,
        "experience": _extract_experience(sections.get("experience", [])),
        "education": _extract_education(sections.get("education", [])),
        "projects": _extract_projects(sections.get("projects", [])),
        "certifications": _simple_items(sections.get("certifications", [])),
        "languages": _simple_items(sections.get("languages", []), split_commas=True),
        "raw_text": cleaned,
    }


def _extension(filename: str) -> str:
    name = (filename or "").strip().lower()
    if "." not in name:
        return ""
    return "." + name.rsplit(".", 1)[-1]


def _extract_text(data: bytes, ext: str, content_type: str) -> tuple[str, int | None]:
    if ext == ".pdf" or content_type == "application/pdf":
        return _extract_pdf(data)
    if ext == ".docx" or content_type.endswith("wordprocessingml.document"):
        return _extract_docx(data), None
    try:
        return data.decode("utf-8"), None
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="ignore"), None


def _extract_pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    chunks = []
    for page in reader.pages:
        chunks.append(page.extract_text() or "")
    return "\n".join(chunks), len(reader.pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(parts)


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ").replace("\u2022", "-").replace("•", "-")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _heading_key(line: str) -> str | None:
    compact = re.sub(r"[\s|:.\-]+", " ", line).strip().lower()
    compact = compact.strip(":-")
    return SECTION_LOOKUP.get(compact)


def _split_sections(lines: list[str]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {}
    current = "header"
    bucket: list[str] = []
    for line in lines:
        key = _heading_key(line)
        if key:
            sections.setdefault(current, [])
            sections[current].extend(bucket)
            current = key
            bucket = []
            continue
        bucket.append(line)
    sections.setdefault(current, [])
    sections[current].extend(bucket)
    return sections


def _first_match(pattern: re.Pattern[str], text: str) -> str | None:
    match = pattern.search(text)
    return match.group(0).rstrip("/") if match else None


def _unique(values) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        if not value:
            continue
        key = value.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(value)
    return out


def _normalize_phone(raw: str) -> str:
    return re.sub(r"\s+", " ", raw).strip(" .-")


def _looks_like_phone(value: str) -> bool:
    digits = re.sub(r"\D", "", value)
    if len(digits) < 10 or len(digits) > 15:
        return False
    if YEAR_RE.fullmatch(digits):
        return False
    return True


def _normalize_url(url: str | None) -> str | None:
    if not url:
        return None
    url = url.rstrip(".,);")
    if url.lower().startswith("www."):
        return "https://" + url
    return url


def _pick_website(text: str, linkedin: str | None, github: str | None) -> str | None:
    skip = {(linkedin or "").lower(), (github or "").lower()}
    for match in URL_RE.findall(text):
        cleaned = _normalize_url(match)
        if not cleaned:
            continue
        lowered = cleaned.lower()
        if any(part and part in lowered for part in skip):
            continue
        if "linkedin.com" in lowered or "github.com" in lowered:
            continue
        if any(token in lowered for token in ("mailto:", "schema.org", "w3.org")):
            continue
        return cleaned
    return None


def _extract_name(lines: list[str], email: str | None, phones: list[str]) -> str | None:
    blocked = {email.lower() if email else ""}
    blocked.update(phone.lower() for phone in phones)
    for line in lines[:10]:
        lowered = line.lower()
        if lowered in blocked or "@" in line:
            continue
        if _looks_like_phone(line) or URL_RE.search(line):
            continue
        if _heading_key(line):
            continue
        words = [word for word in re.split(r"\s+", line) if word]
        if not (1 <= len(words) <= 5) or len(line) > 60:
            continue
        if any(char.isdigit() for char in line):
            continue
        if all(_name_token_ok(word) for word in words):
            if line.isupper():
                return line.title()
            return line
    if email:
        local = email.split("@", 1)[0]
        guess = re.sub(r"[._\-]+", " ", local).strip()
        if guess:
            return guess.title()
    return None


def _name_token_ok(word: str) -> bool:
    cleaned = word.strip(".,'")
    if not cleaned:
        return False
    if cleaned.lower() in {"resume", "cv", "curriculum", "vitae"}:
        return False
    return bool(re.fullmatch(r"[A-Za-z][A-Za-z.'\-]*", cleaned))


def _extract_location(lines: list[str], sections: dict[str, list[str]]) -> str | None:
    header = sections.get("header", lines[:12])
    for line in header[:12]:
        match = CITY_STATE_RE.search(line)
        if match:
            return f"{match.group(1).strip()}, {match.group(2).strip()}"
        lowered = line.lower()
        if lowered in {"remote", "united states", "usa", "u.s.", "u.s.a."}:
            return line
    match = CITY_STATE_RE.search("\n".join(lines[:20]))
    if match:
        return f"{match.group(1).strip()}, {match.group(2).strip()}"
    return None


def _extract_summary(sections: dict[str, list[str]], lines: list[str]) -> str | None:
    block = sections.get("summary") or []
    if not block:
        header = sections.get("header") or lines[:8]
        extras = [line for line in header[1:] if len(line) > 80]
        block = extras[:3]
    text = " ".join(block).strip()
    if len(text) < 40:
        return None
    return re.sub(r"\s+", " ", text)[:1200]


def _extract_skills(skill_lines: list[str], full_text: str) -> list[str]:
    found: list[str] = []
    if skill_lines:
        blob = " ".join(skill_lines)
        parts = re.split(r"[,|/•·\n;]+", blob)
        for part in parts:
            skill = re.sub(r"^[\-–—*]\s*", "", part).strip(" .")
            skill = re.sub(r"^(?:languages|frameworks|tools|databases|cloud|soft skills)\s*[:\-]\s*", "", skill, flags=re.I)
            if 1 < len(skill) <= 40 and not EMAIL_RE.search(skill):
                found.append(_canonical_skill(skill))
    for match in SKILL_PATTERN.findall(URL_RE.sub(" ", full_text)):
        found.append(SKILL_CANONICAL[match.lower()])
    return _unique(found)[:60]


def _canonical_skill(skill: str) -> str:
    return SKILL_CANONICAL.get(skill.lower(), skill)


def _extract_experience(lines: list[str]) -> list[dict[str, Any]]:
    if not lines:
        return []
    blocks = _split_dated_blocks(lines)
    jobs: list[dict[str, Any]] = []
    for block in blocks:
        job = _parse_job_block(block)
        if job:
            jobs.append(job)
    return jobs[:20]


def _split_dated_blocks(lines: list[str]) -> list[list[str]]:
    indexes = [index for index, line in enumerate(lines) if DATE_RANGE_RE.search(line)]
    if not indexes:
        chunks: list[list[str]] = []
        current: list[str] = []
        for line in lines:
            if current and not line.startswith(("-", "*")) and _looks_like_role_line(line):
                chunks.append(current)
                current = [line]
            else:
                current.append(line)
        if current:
            chunks.append(current)
        return chunks

    starts: list[int] = []
    for index in indexes:
        start = index
        if index > 0 and not DATE_RANGE_LINE_RE.match(lines[index]):
            start = index
        else:
            lookback = index - 1
            while lookback >= 0 and lookback >= index - 3:
                if lines[lookback].startswith(("-", "*")):
                    break
                start = lookback
                lookback -= 1
        starts.append(start)

    unique_starts = []
    for start in starts:
        if not unique_starts or start != unique_starts[-1]:
            unique_starts.append(start)

    blocks = []
    for i, start in enumerate(unique_starts):
        end = unique_starts[i + 1] if i + 1 < len(unique_starts) else len(lines)
        block = lines[start:end]
        if block:
            blocks.append(block)
    return blocks


def _looks_like_role_line(line: str) -> bool:
    lowered = line.lower()
    return any(hint in lowered for hint in TITLE_HINTS) or " | " in line


def _parse_job_block(lines: list[str]) -> dict[str, Any] | None:
    bullets = []
    meta = []
    for line in lines:
        if line.startswith(("-", "*")) or (len(line) > 70 and line[:1].islower()):
            bullets.append(re.sub(r"^[\-–—*]\s*", "", line).strip())
        else:
            meta.append(line)
    if not meta:
        return None

    start_date = None
    end_date = None
    is_current = False
    leftover = []
    for line in meta:
        match = DATE_RANGE_RE.search(line)
        if match:
            start_date = _clean_date(match.group("start"))
            end_date = _clean_date(match.group("end"))
            is_current = bool(end_date and end_date.lower() in {"present", "current", "now"})
            remainder = DATE_RANGE_RE.sub("", line).strip(" |-–—,")
            if remainder:
                leftover.append(remainder)
        else:
            leftover.append(line)

    title = None
    company = None
    location = None
    if leftover and " | " in leftover[0]:
        parts = [part.strip() for part in leftover[0].split("|") if part.strip()]
        if len(parts) >= 1:
            title = parts[0]
        if len(parts) >= 2:
            company = parts[1]
        if len(parts) >= 3:
            location = parts[2]
        leftover = leftover[1:]
    elif leftover:
        first = leftover[0]
        second = leftover[1] if len(leftover) > 1 else ""
        if _looks_like_role_line(first):
            title, company = first, second or None
        elif _looks_like_role_line(second):
            company, title = first, second
        else:
            title = first
            company = second or None
        leftover = leftover[2:] if second else leftover[1:]
        if leftover and not location:
            loc_match = CITY_STATE_RE.search(leftover[0])
            if loc_match or leftover[0].lower() in {"remote", "hybrid"}:
                location = leftover[0]
                leftover = leftover[1:]

    if leftover:
        bullets = leftover + bullets
    if not title and not company:
        return None
    return {
        "title": title,
        "company": company,
        "location": location,
        "start_date": start_date,
        "end_date": end_date,
        "is_current": is_current,
        "bullets": bullets[:12],
    }


def _clean_date(value: str | None) -> str | None:
    if not value:
        return None
    return re.sub(r"\s+", " ", value).strip(" .")


def _extract_education(lines: list[str]) -> list[dict[str, Any]]:
    if not lines:
        return []
    blocks = _split_dated_blocks(lines)
    if len(blocks) <= 1 and len(lines) > 4:
        blocks = []
        current: list[str] = []
        for line in lines:
            if current and (DEGREE_RE.search(line) or "university" in line.lower() or "college" in line.lower()):
                blocks.append(current)
                current = [line]
            else:
                current.append(line)
        if current:
            blocks.append(current)

    items = []
    for block in blocks:
        item = _parse_education_block(block)
        if item:
            items.append(item)
    return items[:12]


def _parse_education_block(lines: list[str]) -> dict[str, Any] | None:
    blob = " | ".join(lines)
    degree_match = DEGREE_RE.search(blob)
    gpa_match = GPA_RE.search(blob)
    year_values = YEAR_RE.findall(blob)
    institution = None
    field = None
    degree = degree_match.group(0) if degree_match else None
    for line in lines:
        lowered = line.lower()
        if any(token in lowered for token in ("university", "college", "institute", "school", "polytechnic")):
            institution = re.sub(r"\s+" + DATE_TOKEN, "", line, flags=re.I).strip(" ,|-")
            break
    if not institution:
        for line in lines:
            if not DEGREE_RE.search(line) and not DATE_RANGE_RE.search(line):
                institution = line
                break
    if degree_match:
        after = blob[degree_match.end() :]
        after = re.sub(r"^[\s,|.\-]+(?:in\s+)?", "", after, flags=re.I)
        field_match = re.match(r"([A-Za-z][A-Za-z0-9 &/]+)", after)
        if field_match:
            candidate = field_match.group(1).strip(" .")
            if 2 < len(candidate) < 60 and "university" not in candidate.lower():
                field = candidate
    if not degree and not institution:
        return None
    start_year = year_values[0] if year_values else None
    end_year = year_values[-1] if year_values else None
    return {
        "degree": degree,
        "field": field,
        "institution": institution,
        "start_year": start_year,
        "end_year": end_year,
        "gpa": gpa_match.group(1) if gpa_match else None,
    }


def _extract_projects(lines: list[str]) -> list[dict[str, Any]]:
    if not lines:
        return []
    projects: list[dict[str, Any]] = []
    current_name = None
    bullets: list[str] = []
    for line in lines:
        if line.startswith(("-", "*")):
            bullets.append(re.sub(r"^[\-–—*]\s*", "", line).strip())
            continue
        if current_name:
            projects.append({"name": current_name, "description": " ".join(bullets).strip() or None})
        current_name = line.split("|", 1)[0].strip(" :-")
        bullets = []
    if current_name:
        projects.append({"name": current_name, "description": " ".join(bullets).strip() or None})
    return [project for project in projects if project["name"]][:15]


def _simple_items(lines: list[str], split_commas: bool = False) -> list[str]:
    items: list[str] = []
    for line in lines:
        chunks = re.split(r"[,;|/]+", line) if split_commas else [line]
        for chunk in chunks:
            value = re.sub(r"^[\-–—*]\s*", "", chunk).strip(" .")
            if 1 < len(value) <= 80:
                items.append(value)
    return _unique(items)[:20]
