"""Convert PDF files to DOCX and DOCX files to PDF."""

from __future__ import annotations

import io
import re
from pathlib import Path

from jobscraper.resume_parser import MAX_UPLOAD_BYTES

PDF_TYPE = "application/pdf"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_WINDOWS_FONTS = (
    Path(r"C:\Windows\Fonts\arial.ttf"),
    Path(r"C:\Windows\Fonts\calibri.ttf"),
    Path(r"C:\Windows\Fonts\segoeui.ttf"),
)
_LINUX_FONTS = (
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
)

SECTION_HEADINGS = {
    "education",
    "experience",
    "work experience",
    "professional experience",
    "employment",
    "work history",
    "skills",
    "technical skills",
    "core skills",
    "projects",
    "summary",
    "profile",
    "professional summary",
    "objective",
    "career objective",
    "certifications",
    "certificates",
    "languages",
    "contact",
    "achievements",
    "awards",
    "interests",
    "publications",
    "about me",
    "declaration",
}


class ConvertError(ValueError):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


def pdf_to_docx(data: bytes, filename: str = "") -> bytes:
    _validate_upload(data, filename, expected="pdf")
    from pypdf import PdfReader
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ConvertError("Could not read PDF. The file may be corrupt.", 422) from exc

    if not reader.pages:
        raise ConvertError("PDF has no pages.", 422)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(15, 23, 42)

    extracted = False
    first = True
    for index, page in enumerate(reader.pages):
        text = _page_text(page)
        lines = _split_lines(text)
        if index and lines:
            document.add_page_break()
        if not lines:
            continue
        extracted = True
        for line in lines:
            kind = "h1" if first else _heading_kind(line)
            first = False
            _add_docx_block(document, line, kind)

    if not extracted:
        raise ConvertError(
            "No extractable text in PDF. Image-only / scanned PDFs are not supported (no OCR).",
            422,
        )
    return _save_docx(document)


def docx_to_pdf(data: bytes, filename: str = "") -> bytes:
    _validate_upload(data, filename, expected="docx")
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ConvertError("Could not read DOCX. The file may be corrupt.", 422) from exc

    blocks = _blocks_from_docx(document)
    if not any(text for text, kind in blocks if kind != "blank"):
        raise ConvertError("DOCX has no extractable text.", 422)

    if blocks and blocks[0][1] == "p":
        blocks[0] = (blocks[0][0], "h1")

    try:
        return _render_pdf(blocks)
    except ConvertError:
        raise
    except Exception as exc:
        try:
            return _fallback_pdf(blocks)
        except Exception:
            raise ConvertError(f"Could not convert DOCX to PDF: {exc}", 422) from exc


def _page_text(page) -> str:
    try:
        text = page.extract_text(extraction_mode="layout") or ""
    except Exception:
        text = ""
    if len(text.strip()) < 8:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
    return text


def _add_docx_block(document, text: str, kind: str) -> None:
    from docx.shared import Pt

    if kind == "h1":
        document.add_heading(text, level=0)
        return
    if kind == "h2":
        document.add_heading(text.rstrip(":"), level=1)
        return
    if kind == "li":
        document.add_paragraph(_strip_bullet(text), style="List Bullet")
        return
    paragraph = document.add_paragraph(text)
    if paragraph.runs:
        paragraph.runs[0].font.size = Pt(11)


def _blocks_from_docx(document) -> list[tuple[str, str]]:
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    blocks: list[tuple[str, str]] = []
    seen: set[str] = set()

    def add_text(text: str, style_name: str = "") -> None:
        text = (text or "").strip()
        if not text:
            if blocks and blocks[-1][1] != "blank":
                blocks.append(("", "blank"))
            return
        key = re.sub(r"\s+", " ", text).lower()
        if key in seen:
            return
        seen.add(key)
        kind = "li" if _strip_bullet(text) != text else _heading_kind(text, style_name)
        blocks.append((text, kind))

    def add_paragraph(paragraph) -> None:
        try:
            style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
        except Exception:
            style_name = ""
        add_text(paragraph.text or "", style_name)

    def add_table(table) -> None:
        try:
            rows = table.rows
        except Exception:
            return
        for row in rows:
            for cell in _unique_cells(row):
                for paragraph in cell.paragraphs:
                    add_paragraph(paragraph)
                try:
                    nested = cell.tables
                except Exception:
                    nested = []
                for nested_table in nested:
                    add_table(nested_table)

    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            add_paragraph(Paragraph(child, document))
        elif tag == "tbl":
            add_table(Table(child, document))

    for txbx in document.element.iter(qn("w:txbxContent")):
        for para in txbx.iter(qn("w:p")):
            text = "".join(node.text or "" for node in para.iter(qn("w:t"))).strip()
            add_text(text)

    for section in document.sections:
        for part in (section.header, section.footer):
            try:
                for paragraph in part.paragraphs:
                    add_paragraph(paragraph)
            except Exception:
                continue

    return blocks


def _unique_cells(row):
    seen = []
    for cell in row.cells:
        ident = id(cell._tc)
        if ident in seen:
            continue
        seen.append(ident)
        yield cell


def _heading_kind(text: str, style_name: str = "") -> str:
    style_name = (style_name or "").lower()
    if "heading 1" in style_name or style_name == "title":
        return "h1"
    if "heading" in style_name:
        return "h2"
    compact = re.sub(r"[\s|:.\-]+$", "", text or "").strip()
    if compact.lower() in SECTION_HEADINGS:
        return "h2"
    letters = re.sub(r"[^A-Za-z ]", "", compact)
    if letters.isupper() and 3 <= len(letters) <= 36:
        return "h2"
    if _strip_bullet(text) != (text or "").strip():
        return "li"
    return "p"


def _strip_bullet(text: str) -> str:
    return re.sub(r"^[\-•●○▪◦\*]+\s*", "", (text or "").strip())


def preview_text_from_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = []
    for text, kind in _blocks_from_docx(document):
        if kind == "blank":
            if lines and lines[-1] != "":
                lines.append("")
            continue
        lines.append(text)
    return "\n".join(lines).strip()


def _validate_upload(data: bytes, filename: str, expected: str) -> None:
    if not data:
        raise ConvertError("Empty file.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ConvertError("File is larger than 8 MB.", 413)

    ext = _extension(filename)
    if expected == "pdf":
        if data.startswith(b"%PDF"):
            return
        if ext == ".doc":
            raise ConvertError("Old .doc is not supported. Save as PDF or DOCX first.", 415)
        raise ConvertError("Upload a PDF file (.pdf).", 415)
    if expected == "docx":
        if _looks_like_docx(data):
            return
        if ext == ".doc":
            raise ConvertError("Old .doc is not supported. Save as DOCX or convert to PDF first.", 415)
        raise ConvertError("Upload a Word file (.docx).", 415)


def _looks_like_docx(data: bytes) -> bool:
    return data[:2] == b"PK"


def _extension(filename: str) -> str:
    name = (filename or "").strip().lower()
    if "." not in name:
        return ""
    return "." + name.rsplit(".", 1)[-1]


def _split_lines(text: str) -> list[str]:
    cleaned = (text or "").replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    cleaned = cleaned.replace("\u00a0", " ")
    lines = []
    for raw in cleaned.split("\n"):
        line = re.sub(r"[ \t]+", " ", raw).strip()
        if line:
            lines.append(line)
    return lines


def _save_docx(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _font_file() -> Path | None:
    for path in (*_WINDOWS_FONTS, *_LINUX_FONTS):
        if path.exists():
            return path
    return None


def _render_pdf(blocks: list[tuple[str, str]]) -> bytes:
    try:
        from fpdf import FPDF
    except ImportError:
        return _fallback_pdf(blocks)

    pdf = FPDF(format="Letter", unit="mm")
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    pdf.set_margins(16, 16, 16)

    font_path = _font_file()
    family = "Helvetica"
    if font_path is not None:
        try:
            family = "Body"
            pdf.add_font("Body", "", str(font_path))
            bold_path = _bold_companion(font_path)
            pdf.add_font("Body", "B", str(bold_path or font_path))
        except Exception:
            family = "Helvetica"
            font_path = None

    def write(text: str, size: int, style: str = "", space_after: float = 3, color=(15, 23, 42)) -> None:
        payload = _pdf_text(text)
        if font_path is None:
            payload = _latin1(payload)
        pdf.set_text_color(*color)
        try:
            pdf.set_font(family, style=style, size=size)
            pdf.multi_cell(0, max(size * 0.42 + 2, 5), payload)
        except Exception:
            pdf.set_font("Helvetica", size=size)
            pdf.multi_cell(0, 6, _latin1(payload))
        pdf.ln(space_after)

    for text, kind in blocks:
        if kind == "blank":
            pdf.ln(2)
            continue
        if kind == "h1":
            write(text, 20, "B", 2)
            continue
        if kind == "h2":
            pdf.ln(2)
            write(text.rstrip(":").upper(), 12, "B", 1, (37, 99, 235))
            pdf.set_draw_color(37, 99, 235)
            pdf.set_line_width(0.3)
            y = pdf.get_y()
            pdf.line(16, y, 216 - 16, y)
            pdf.ln(2)
            continue
        if kind == "li":
            write(f"- {_strip_bullet(text)}", 11, "", 1.5)
            continue
        write(text, 11, "", 1.5)

    output = pdf.output()
    return bytes(output)


def _pdf_text(text: str) -> str:
    cleaned = "".join(ch if ch == "\n" or ord(ch) >= 32 else " " for ch in (text or ""))
    cleaned = cleaned.replace("\x00", " ").strip()
    parts = []
    for token in cleaned.split(" "):
        if len(token) > 90:
            parts.extend(token[i : i + 90] for i in range(0, len(token), 90))
        else:
            parts.append(token)
    return " ".join(parts)


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _fallback_pdf(blocks: list[tuple[str, str]]) -> bytes:
    lines: list[str] = []
    for text, kind in blocks:
        if kind == "blank":
            lines.append("")
            continue
        payload = _latin1(_pdf_text(text))
        if kind == "h2":
            payload = payload.upper()
        if kind == "li":
            payload = "- " + _latin1(_strip_bullet(text))
        while len(payload) > 90:
            lines.append(payload[:90])
            payload = payload[90:]
        lines.append(payload)
    if not lines:
        lines = [" "]

    per_page = 50
    pages = [lines[i : i + per_page] for i in range(0, len(lines), per_page)]
    content_streams: list[bytes] = []
    for page_lines in pages:
        cmds = ["BT", "/F1 11 Tf", "50 742 Td"]
        for index, line in enumerate(page_lines):
            esc = _pdf_escape(line)
            cmds.append(f"({esc}) Tj" if index == 0 else f"0 -14 Td ({esc}) Tj")
        cmds.append("ET")
        content_streams.append("\n".join(cmds).encode("latin-1", "replace"))

    n_pages = len(pages)
    page_ids = list(range(4, 4 + n_pages))
    content_ids = list(range(4 + n_pages, 4 + 2 * n_pages))
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    kids = " ".join(f"{item} 0 R" for item in page_ids)
    objects[2] = f"<< /Type /Pages /Count {n_pages} /Kids [{kids}] >>".encode("ascii")
    for page_id, content_id, stream in zip(page_ids, content_ids, content_streams):
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        ).encode("ascii")
        objects[content_id] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode("ascii") + stream + b"\nendstream"
        )

    max_id = max(objects)
    out = bytearray(b"%PDF-1.4\n")
    offsets = {0: 0}
    for obj_id in range(1, max_id + 1):
        offsets[obj_id] = len(out)
        out += f"{obj_id} 0 obj\n".encode("ascii") + objects[obj_id] + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {max_id + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for obj_id in range(1, max_id + 1):
        out += f"{offsets[obj_id]:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {max_id + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


def _bold_companion(regular: Path) -> Path | None:
    name = regular.name.lower()
    mapping = {
        "arial.ttf": "arialbd.ttf",
        "calibri.ttf": "calibrib.ttf",
        "segoeui.ttf": "segoeuib.ttf",
        "dejavusans.ttf": "DejaVuSans-Bold.ttf",
        "liberationsans-regular.ttf": "LiberationSans-Bold.ttf",
    }
    bold_name = mapping.get(name)
    if not bold_name:
        return None
    candidate = regular.with_name(bold_name)
    return candidate if candidate.exists() else None


def _latin1(text: str) -> str:
    return text.encode("latin-1", "replace").decode("latin-1")
